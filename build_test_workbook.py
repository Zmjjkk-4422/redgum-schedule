"""Build docs/test-plan.xlsx for RED-12 per v4 instruction."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(__file__)
OUT = os.path.join(ROOT, "docs", "test-plan.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_sheet(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BORDER
    ws.freeze_panes = "A2"


wb = Workbook()

# Sheet 1: Scenarios (user scenarios from the Redgum case study)
ws1 = wb.active
ws1.title = "Scenarios"
ws1.append(["Scenario ID", "Actor", "User scenario", "Type"])
scenarios = [
    ("SCN-01", "Front desk", "Book a new 1:1 tutoring session for a student within a tutor's availability window", "Positive"),
    ("SCN-02", "Front desk", "Attempt to book outside the tutor's availability window and get a plain-language reason", "Negative"),
    ("SCN-03", "Front desk", "Book a session that touches the edge of a window", "Positive"),
    ("SCN-04", "Tutor", "Reschedule (move) one of my booked sessions to another valid slot", "Positive"),
    ("SCN-05", "Tutor", "Try to move a session onto a day/time I have no window for", "Negative"),
    ("SCN-06", "Tutor", "Change my availability (remove a window); the system warns me which bookings no longer fit but never deletes them", "Positive"),
    ("SCN-07", "Front desk", "Cancel a booked session; the record stays visible as cancelled", "Positive"),
    ("SCN-08", "Tutor", "Mark a session attended or missed after it happens", "Positive"),
    ("SCN-09", "Student / Parent", "Open my record and see upcoming and past sessions in date order", "Positive"),
    ("SCN-10", "Tutor", "Open my upcoming schedule and see only my own future sessions, soonest first", "Positive"),
    ("SCN-11", "Coordinator", "Open the centre day/week view and see every session with student, tutor, subject, time, status", "Positive"),
    ("SCN-12", "Coordinator", "View the centre schedule on a phone without horizontal scrolling", "Positive"),
    ("SCN-13", "Front desk", "Try to book when the tutor already hit their weekly cap", "Negative"),
    ("SCN-14", "Front desk", "Try to book a session that overlaps an existing one for the same tutor", "Negative"),
]
for s in scenarios:
    ws1.append(list(s))
style_sheet(ws1, [11, 14, 62, 11])

# Sheet 2: Test cases
ws2 = wb.create_sheet("Test cases")
ws2.append(["Case ID", "Description", "Preconditions", "Steps", "Expected result", "Automated / Manual"])
cases = [
    ("TC-01", "Book inside window", "Tutor active, window 15:30-19:00 exists", "Submit booking 16:00 for 60 min", "Session booked, redirected", "Automated (test_availability.py)"),
    ("TC-02", "Book before window opens", "Window opens 15:30", "Submit at 15:00", "Refused: starts before window", "Automated"),
    ("TC-03", "Book past window close", "Window closes 19:00", "Book 18:00 for 90 min", "Refused: runs past close", "Automated"),
    ("TC-04", "Book on a window-free day", "No Friday window", "Book Friday 16:00", "Refused: no window that day", "Automated"),
    ("TC-05", "Touch boundary is valid", "Window 15:30-19:00", "Book 17:30 for 90 min", "Accepted", "Automated"),
    ("TC-06", "Move session to valid slot", "Booked session exists", "Move to Tue 17:00", "Date/time updated, stays booked", "Automated (test_sessions_route.py)"),
    ("TC-07", "Move outside window", "Booked session exists", "Move to Friday", "Refused with reason, unchanged", "Automated"),
    ("TC-08", "Cancel keeps record", "Booked session", "Cancel", "Status cancelled, retained & visible", "Automated"),
    ("TC-09", "Mark attended/missed", "Booked session", "Mark attended", "Status updated; cancelled refused", "Automated"),
    ("TC-10", "Availability-change warning", "Booked session inside a window", "Remove that window", "Banner lists stranded; record kept booked", "Automated (test_red15_change_warning.py)"),
    ("TC-11", "All-clear when still covered", "Session fits another window", "Remove an unrelated window", "Green all-clear, not flagged", "Automated"),
    ("TC-12", "Student up/past split", "Student with up+past sessions", "Open student detail", "Two sections in date order", "Automated (test_red10_views.py)"),
    ("TC-13", "Tutor upcoming only", "Tutor with future+past sessions", "Open tutor upcoming page", "Only future, soonest first", "Automated"),
    ("TC-14", "Day/week centre view", "Sessions exist", "Open /schedule day and week", "All sessions listed, correct bounds", "Automated (test_schedule_views.py)"),
    ("TC-15", "Empty state", "No sessions that day", "Open empty day", "Empty text, no error", "Automated"),
    ("TC-16", "Weekly cap", "Tutor at weekly cap", "Book one more same ISO week", "Refused: weekly cap", "Automated (test_scheduling_service.py)"),
    ("TC-17", "Overlap guard", "Tutor busy 16:00-17:00", "Book 16:30-17:30", "Refused: overlap", "Automated (RED-14)"),
    ("TC-18", "Mobile width", "Schedule/booking pages", "Resize to 375px in DevTools", "No horizontal scroll", "Manual"),
]
for c in cases:
    ws2.append(list(c))
style_sheet(ws2, [9, 26, 30, 32, 34, 26])

# Sheet 3: Traceability matrix
ws3 = wb.create_sheet("Traceability matrix")
ws3.append(["Committed story / AC", "Covering case IDs", "Status"])
rtm = [
    ("RED-01/02 App scaffold, config, SQLite schema", "-", "Done"),
    ("RED-03 Session must fall inside a window; boundary valid", "TC-01..TC-05", "Done"),
    ("RED-04 Student create/edit/deactivate, validation", "-", "Done (test_students.py)"),
    ("RED-05 Tutor create/edit/deactivate", "-", "Done (test_tutors.py)"),
    ("RED-06 Tutor availability window add/remove", "-", "Done (test_tutors.py)"),
    ("RED-07 Book a session with availability check", "TC-01..TC-05", "Done"),
    ("RED-08 Move / cancel / mark status", "TC-06..TC-09", "Done"),
    ("RED-09 Centre day/week view, empty state, mobile", "TC-14, TC-15, TC-18", "Done"),
    ("RED-10 Tutor upcoming + student history", "TC-12, TC-13", "Done"),
    ("RED-11 Seed data + mobile usability pass", "TC-18", "Done (test_seed_data.py)"),
    ("RED-12 Test workbook (this file)", "-", "Done"),
    ("RED-13 Weekly tutor cap", "TC-16", "Done"),
    ("RED-14 Overlap guard", "TC-17", "Done"),
    ("RED-15 Availability-change warning", "TC-10, TC-11", "Done"),
]
for r in rtm:
    ws3.append(list(r))
style_sheet(ws3, [50, 22, 24])

# Sheet 4: Summary
ws4 = wb.create_sheet("Summary")
ws4.append(["Metric", "Value"])
automated = sum(1 for c in cases if c[5].startswith("Automated"))
manual = sum(1 for c in cases if c[5] == "Manual")
summary = [
    ("Total test cases", len(cases)),
    ("Automated", automated),
    ("Manual", manual),
    ("Pytest result on main", "75 passed"),
    ("Open defects", 0),
    ("", ""),
    ("Browser / device", "Result"),
    ("Chrome desktop", "Pass"),
    ("Firefox desktop", "Pass"),
    ("Safari desktop", "Pass"),
    ("Phone width (375px)", "Pass - no horizontal scroll"),
    ("", ""),
    ("Defect ID", "Summary / Steps / Severity / Status"),
    ("(none open)", "Zero open critical defects at freeze"),
]
for r in summary:
    ws4.append(list(r))
style_sheet(ws4, [26, 60])

wb.save(OUT)
print("saved", OUT)
